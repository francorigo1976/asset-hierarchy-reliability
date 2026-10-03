"""Exact SAP PM LSMW workbook export — port of the app's exportSAPPM.

Layout matches the client's "FL Upload / EQ Upload - New Template" files: SAP field codes on Excel row 7,
lengths row 8, descriptions row 9, example row 14, real data from row 17. Over-length cells are filled red with a
comment and listed on the trailing "Validation Summary" sheet; labelled fields are truncated to SAP limits and logged.
"""
from __future__ import annotations

import re
from datetime import date
from pathlib import Path
from typing import Optional

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .classify import ZCM_DESC
from .sapname import generate_sap_name
from .hierarchy import FLOC_LABEL_LIMIT, FLOC_SEGMENT_MAX, Hierarchy
from .models import Node

FL_COLS, EQ_COLS, DATA_ROW0 = 56, 101, 16     # 0-indexed array position of Excel row 17

# (field code, length, description) for columns 1..N (column 0 holds the row label)
FL_SPEC = [
    ("TPLNR", 30, "Functional location"), ("TPLKZ", 5, "Structure indicator"), ("FLTYP", 1, "Functional location category"),
    ("TRPNR", 30, "Reference functional location"), ("PLTXT", 40, "Functional location description"),
    ("SWERK", 4, "Maintenance plant"), ("STORT", 10, "Location"), ("MSGRP", 8, "Room"), ("BEBER", 3, "Plant section"),
    ("ARBPL", 8, "Work center"), ("ABCKZ", 1, "ABC indicator"), ("EQFNR", 30, "Sort field"),
    ("BUKRS", 4, "Compagny code"), ("ANLNR", 12, "Main asset number"), ("ANLUN", 4, "Asset sub-number"),
    ("GSBER", 4, "Business Area"), ("KOSTL", 10, "Cost Center"), ("EQART", 10, "Object Type"), ("PROID", 24, "WBS Element"),
    ("DAUFN", 12, "Standing order number"), ("TPLMA", 30, "Superior functional location"),
    ("SUBMT", 18, "Construction type"), ("IEQUI", 1, "Equipment installation allowed"), ("EINZL", 1, "Single installation"),
    ("IWERK", 4, "Maintenance Planning Plant"), ("INGRP", 3, "Planner Group"), ("GEWRK", 8, "Main work center"),
    ("WERGW", 4, "Main work center plant"), ("RBNR", 9, "Catalog Profile"), ("BEGRU", 4, "Authorization group"),
    ("VKORG", 4, "Sales organization"), ("VTWEG", 2, "Distribution channel"), ("SPART", 2, "Division"),
    ("POSNR", 4, "Position"), ("AUFNR", 12, "Settlement order"), ("DATAB", 8, "Start-up Date of the Technical Object"),
    ("STRNO", 40, "Functional location label"), ("STRNO_TPLMA", 40, "Functional location label"),
    ("EQART", 10, "Object type"), ("INVNR", 25, "Inventory number"), ("BRGEW", 17, "Gross Weight"),
    ("GEWEI", 3, "Weight Unit"), ("GROES", 18, "Size/dimension"), ("ANSWT", 17, "AcquisitionValue: IBIP Character Structure"),
    ("WAERS", 5, "Currency Key"), ("ANSDT", 8, "Acquisition date"), ("HERST", 30, "Manufacturer of asset"),
    ("HERLD", 3, "Country of manufacture"), ("TYPBZ", 20, "Manufacturer model number"), ("BAUJJ", 4, "Year of construction"),
    ("BAUMM", 2, "Month of construction"), ("MAPAR", 30, "Manufacturer part number"),
    ("SERGE", 30, "Manufacturer serial number"), ("VKBUR", 4, "Sales office"), ("VKGRP", 3, "Sales group"),
]
_EQ_CODES = ("LEGACYKEY EQUNR DATSL EQTYP EQKTX BEGRU EQART GROES INVNR BRGEW GEWEI ELIEF ANSDT ANSWT WAERS HERST HERLD "
             "BAUJJ BAUMM TYPBZ SERGE MAPAR GERNR GWLEN KUND1 KUND2 KUND3 SWERK STORT MSGRP BEBER ARBPL ABCKZ EQFNR BUKRS "
             "ANLNR ANLUN GSBER KOSTL PROID DAUFN AUFNR TIDNR SUBMT HEQUI HEQNR EINZL IWERK INGRP GEWRK WERGW RBNR TPLNR "
             "DISMANTLE VKORG VTWEG SPART MATNR SERNR WERK LAGER CHARGE KUNDE KZKBL PLANV FGRU1 FGRU2 STEUF STEUF_REF "
             "KTSCH KTSCH_REF EWFORM EWFORM_REF BZOFFB BZOFFB_REF OFFSTB EHOFFB OFFSTB_REF BZOFFE BZOFFE_REF OFFSTE EHOFFE "
             "OFFSTE_REF WARPL IMRC_POINT INDAT INTIM INBDT GWLDT AULDT LIZNR MGANR REFMA VKBUR VKGRP WARR_INBD WAGET "
             "GAERB ACT_CHANGE_AA STRNO").split()
_EQ_LEN = [40, 18, 8, 1, 40, 4, 10, 18, 25, 17, 3, 10, 8, 17, 5, 30, 3, 4, 2, 20, 30, 30, 18, 8, 10, 10, 10, 4, 10, 8, 3, 8,
           1, 30, 4, 12, 4, 4, 10, 24, 12, 12, 25, 18, 18, 4, 1, 4, 3, 8, 4, 9, 30, 1, 4, 2, 2, 18, 18, 4, 4, 10, 10, 1, 3,
           4, 4, 4, 1, 7, 1, 6, 1, 2, 1, 7, 3, 1, 2, 1, 7, 3, 1, 12, 12, 8, 6, 8, 8, 8, 20, 20, 18, 4, 3, 1, 1, 1, 1, 40]
_EQ_DESC = {1: "Old equipment number (key)", 2: "Equipment number", 3: "Validity date of the technical object",
            4: "Equipment category", 5: "Description of the equipment", 6: "Authorization group", 7: "Object type",
            8: "Size/dimension", 9: "Inventory number", 10: "Gross Weight", 11: "Weight Unit", 12: "Vendor number",
            13: "Acquisition date", 14: "AcquisitionValue", 15: "Currency Key", 16: "Manufacturer",
            17: "Country of manufacture", 18: "Year of construction", 19: "Month of construction",
            20: "Description type", 21: "Supplier serial number", 22: "Manufacturer part number", 23: "Serial number",
            28: "Maintenance plant", 29: "Location", 30: "Room", 31: "Plant section", 32: "Work center",
            33: "ABC indicator", 34: "Sort field", 35: "Compagny code", 36: "Main asset number",
            37: "Asset sub-number", 38: "Business Area", 39: "Cost Center", 40: "WBS Element",
            41: "Standing order number", 42: "Settlement order", 43: "Technical identification number",
            44: "Construction type", 45: "Superior Equipment", 46: "Equipment position at install location",
            47: "Single equipment installation at FL", 48: "Maintenance Planning Plant", 49: "Planner Group",
            50: "Main work center", 51: "Plant Main work center", 52: "Catalog Profile", 53: "Functional location",
            88: "First start-up date", 100: "Functional location label"}
assert len(_EQ_CODES) == len(_EQ_LEN) == 100
EQ_SPEC = [(c, _EQ_LEN[i], _EQ_DESC.get(i + 1, "")) for i, c in enumerate(_EQ_CODES)]
FL_IDX = {}
for _i, (_c, _, _) in enumerate(FL_SPEC, start=1):
    FL_IDX.setdefault(_c, _i)           # first occurrence wins; EQART handled explicitly (col 39)
FL_IDX_EQART2 = 39
EQ_IDX = {c: i for i, (c, _, _) in enumerate(EQ_SPEC, start=1)}

BAD_FILL = PatternFill("solid", fgColor="FFC7CE")
BAD_FONT = Font(color="9C0006")
HDR_FILL = PatternFill("solid", fgColor="DCE6F1")
WARN_FILL = PatternFill("solid", fgColor="FFEB9C")


def _row_labels(spec, cols):
    codes, lens, descs = [""] * cols, [""] * cols, [""] * cols
    codes[0], lens[0], descs[0] = "Field", "Length", "Description"
    for i, (c, ln, d) in enumerate(spec, start=1):
        codes[i], lens[i], descs[i] = c, ln, d
    return codes, lens, descs


def _preamble(site, cols, spec, example):
    codes, lens, descs = _row_labels(spec, cols)
    empty = lambda: [""] * cols           # noqa: E731
    site_row, fr = empty(), empty()
    site_row[0], fr[0] = site or "", "description"
    return [empty(), empty(), empty(), empty(), empty(), site_row, codes, lens, descs, fr,
            empty(), empty(), empty(), example or empty(), empty(), empty()]


def _obj_type(node: Node, fallback: str, area: bool = False) -> str:
    if area:
        return "AREA"
    desc = ZCM_DESC.get(node.catalog_code, "") or fallback
    return re.sub(r"[^A-Z0-9]", "_", desc.upper())[:10]


def export_lsmw(h: Hierarchy, path: str | Path, today: Optional[str] = None) -> dict:
    p = h.project
    today = today or date.today().strftime("%Y%m%d")
    truncations: list[dict] = []
    warnings: list[str] = []

    def trunc(val, mx, field=None):
        if val is None or val == "":
            return val
        s = str(val).strip()
        if mx and len(s) > mx:
            if field:
                truncations.append({"field": field, "original": s, "length": len(s), "limit": mx})
            return s[:mx]
        return s

    def year(y):
        s = str(y or "").strip()
        return s if re.fullmatch(r"\d{4}", s) else ""

    def posnr(n: Node, counters: dict) -> str:
        if n.sort_order and n.sort_order > 0:
            return f"{n.sort_order:04d}"
        k = n.parent_id or "__root__"
        counters[k] = counters.get(k, 0) + 10
        return f"{counters[k]:04d}"

    ordered = h.ordered()
    mode = p.sap_structure or "floc-only"
    kids = lambda n: bool(h.children(n.id))                       # noqa: E731
    floc_nodes = [n for n in ordered if n.is_area or mode == "one-to-one"
                  or (mode == "structure-leaf" and not kids(n)) or n.is_also_equipment]
    equip_nodes = [n for n in ordered if (n.is_also_equipment if n.is_area else
                   (kids(n) or n.is_also_equipment) if mode == "structure-leaf" else True)]

    def floc_label(n: Node) -> str:
        lab = h.floc_label(n)
        if len(lab) > FLOC_LABEL_LIMIT:
            warnings.append(f"FLOC label exceeds {FLOC_LABEL_LIMIT} chars: \"{lab}\" ({len(lab)})")
        return lab

    tplkz = trunc(p.sap_structure_indicator if p.sap_structure_indicator and p.sap_structure_indicator != p.prefix
                  else "ZRCS1", 5)

    # ---- New Functional Locations ----
    fl_ex = [""] * FL_COLS
    fl_ex[0], fl_ex[1], fl_ex[2], fl_ex[3], fl_ex[5] = "Example", "EXAMPLE-AREA-01", "ZRCS1", "L", "Example Functional Location"
    fl_ex[18], fl_ex[23], fl_ex[34] = "AREA", "X", "0010"
    fl_data = _preamble(p.site, FL_COLS, FL_SPEC, fl_ex)
    counters: dict = {}
    for n in floc_nodes:
        parent = h.get(n.parent_id)
        r = [""] * FL_COLS
        r[1], r[2], r[3] = floc_label(n), tplkz, "L"
        r[5] = trunc(n.description or n.name, 40, "PLTXT")
        r[6] = trunc(p.sap_maint_plant, 4)
        r[10] = trunc(p.sap_work_center, 8)
        r[12] = trunc(n.tag_number, 30, "EQFNR")
        r[13] = trunc(p.sap_company_code, 4)
        r[17] = trunc(p.sap_cost_center, 10)
        r[21] = floc_label(parent) if parent else ""
        r[23] = "X"
        r[25], r[26], r[27], r[28] = (trunc(p.sap_planning_plant, 4), trunc(p.sap_planner_group, 3),
                                      trunc(p.sap_work_center, 8), trunc(p.sap_maint_plant, 4))
        r[34], r[36] = posnr(n, counters), today
        r[39] = trunc(_obj_type(n, "SYSTEM", area=n.is_area), 10)
        r[47], r[49] = trunc(n.manufacturer, 30, "HERST"), trunc(n.model, 20, "TYPBZ")
        r[50], r[52], r[53] = year(n.year), trunc(n.tann_part_no, 30, "MAPAR"), trunc(n.serial, 30, "SERGE")
        fl_data.append(r)
    fl_chg_note = [""] * FL_COLS
    fl_chg_note[0] = "Only populate cells you want to change"
    fl_chg = _preamble(p.site, FL_COLS, FL_SPEC, fl_chg_note)

    # ---- New Equipment ----
    eq_ex = [""] * EQ_COLS
    eq_ex[0], eq_ex[1], eq_ex[3], eq_ex[4], eq_ex[5], eq_ex[7] = "Example", "EXAMPLE/REF-001", "20240101", "E", "Example Equipment", "EXAMPLE"
    eq_ex[34], eq_ex[53], eq_ex[88] = "EX-001", "EXAMPLE-AREA-01", "20240101"
    eq_data = _preamble(p.site, EQ_COLS, EQ_SPEC, eq_ex)
    for n in equip_nodes:
        install = n if (n.is_area and n.is_also_equipment) else h.get(n.parent_id)
        r = [""] * EQ_COLS
        r[1] = trunc(n.asset_id, 40, "LEGACYKEY")
        r[3], r[4] = today, "E"
        r[5] = trunc(n.sap_name or generate_sap_name(n).sap_name or n.description or n.name, 40, "EQKTX")
        r[7] = trunc(_obj_type(n, "EQUIPMENT"), 10)
        r[16], r[18] = trunc(n.manufacturer, 30, "HERST"), year(n.year)
        r[20], r[21], r[22] = trunc(n.model, 20, "TYPBZ"), trunc(n.serial, 30, "SERGE"), trunc(n.tann_part_no, 30, "MAPAR")
        r[28], r[32] = trunc(p.sap_maint_plant, 4), trunc(p.sap_work_center, 8)
        r[34], r[35], r[39] = trunc(n.tag_number, 30, "EQFNR"), trunc(p.sap_company_code, 4), trunc(p.sap_cost_center, 10)
        r[43] = trunc(n.tag_number, 25)
        parent = h.get(n.parent_id)
        if not n.is_area and parent is not None and not parent.is_area:
            heq = str(parent.asset_id or "").strip()          # never truncated: would orphan the HEQUI lookup
            if len(heq) > 18:
                warnings.append(f"Parent equipment ID exceeds 18 chars (len {len(heq)}): {heq}")
            r[45] = heq
        r[48], r[49], r[50], r[51] = (trunc(p.sap_planning_plant, 4), trunc(p.sap_planner_group, 3),
                                      trunc(p.sap_work_center, 8), trunc(p.sap_maint_plant, 4))
        r[52] = trunc(n.catalog_code, 9, "RBNR")
        flv = floc_label(install).strip() if install else ""
        if len(flv) > 30:
            warnings.append(f"Install FLOC exceeds 30 chars (len {len(flv)}): {flv}")
        r[53], r[88] = flv, today
        eq_data.append(r)
    eq_chg_note = [""] * EQ_COLS
    eq_chg_note[0] = "Only populate cells you want to change"
    eq_chg = _preamble(p.site, EQ_COLS, EQ_SPEC, eq_chg_note)

    # ---- workbook ----
    wb = Workbook()
    wb.remove(wb.active)
    violations: list[dict] = []

    def add_sheet(name, data, lens=None, check=False):
        ws = wb.create_sheet(name[:31])
        for row in data:
            ws.append(list(row))
        for c in range(1, max(len(r) for r in data) + 1):
            width = max(6, max(len(str(r[c - 1])) for r in data if c - 1 < len(r))) + 2
            ws.column_dimensions[get_column_letter(c)].width = min(width, 35)
        if lens is not None:
            for c in ws[7]:
                c.font, c.fill = Font(bold=True), HDR_FILL
        if check:
            for r in range(DATA_ROW0, len(data)):
                for c in range(1, len(data[r])):
                    limit = lens[c] if isinstance(lens[c], int) else 0
                    v = data[r][c]
                    if limit and v not in (None, "") and len(str(v)) > limit:
                        cell = ws.cell(row=r + 1, column=c + 1)
                        cell.fill, cell.font = BAD_FILL, BAD_FONT
                        cell.comment = Comment(f"Exceeds SAP limit: {len(str(v))}/{limit} chars", "HierarchyCapture")
                        violations.append({"sheet": name, "row": r + 1, "col": get_column_letter(c + 1),
                                           "field": data[6][c], "value": str(v), "length": len(str(v)), "limit": limit})
        return ws

    fl_lens, eq_lens = _row_labels(FL_SPEC, FL_COLS)[1], _row_labels(EQ_SPEC, EQ_COLS)[1]
    add_sheet("New Functional Locations", fl_data, fl_lens, True)
    add_sheet("Change Existing Functional Loc", fl_chg, fl_lens)
    add_sheet("New Equipment", eq_data, eq_lens, True)
    add_sheet("Change Existing Equipment", eq_chg, eq_lens)
    add_sheet("To be deleted", [["Functional Location / Equipment No"]])

    def data_cells(data, lens):
        return sum(1 for r in data[DATA_ROW0:] for c in range(1, len(r))
                   if isinstance(lens[c], int) and lens[c] > 0 and r[c] not in (None, ""))
    checked = data_cells(fl_data, fl_lens) + data_cells(eq_data, eq_lens)

    summary = wb.create_sheet("Validation Summary")
    if not violations and not truncations:
        summary.append(["All fields within SAP character limits ✓"])
        summary.append([f"{checked} data fields checked, 0 exceed limits, 0 truncated"])
        summary["A1"].font = Font(bold=True, size=14, color="006100")
    else:
        bits = []
        if violations:
            bits.append(f"{len(violations)} hierarchy field(s) exceed SAP cap")
        if truncations:
            bits.append(f"{len(truncations)} field(s) truncated to SAP limits")
        summary.append([" · ".join(bits)])
        summary["A1"].font = Font(bold=True, size=14, color="9C0006" if violations else "9C5700")
        summary.append([])
        if violations:
            summary.append(["Sheet", "Row", "Column", "Field Name", "Value", "Length", "Limit", "Status"])
            for c in summary[summary.max_row]:
                c.font, c.fill = Font(bold=True), HDR_FILL
            for v in violations:
                summary.append([v["sheet"], v["row"], v["col"], v["field"], v["value"], v["length"], v["limit"], "EXCEEDS LIMIT"])
                for c in summary[summary.max_row]:
                    c.fill, c.font = BAD_FILL, BAD_FONT
            summary.append([])
        if truncations:
            summary.append(["Field", "Original Value", "Original Length", "Limit", "Status"])
            for c in summary[summary.max_row]:
                c.font, c.fill = Font(bold=True), HDR_FILL
            for t in truncations:
                summary.append([t["field"], t["original"], t["length"], t["limit"], "TRUNCATED"])
                for c in summary[summary.max_row]:
                    c.fill = WARN_FILL
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return {"path": path, "flocs": len(floc_nodes), "equipment": len(equip_nodes), "violations": violations,
            "truncations": truncations, "warnings": warnings}
