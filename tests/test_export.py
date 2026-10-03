import pytest
from openpyxl import load_workbook

from data_export import AuditTrail, export_records, get_template, list_templates, validate_records
from data_export.exporter import ExportBlocked

GOOD = [
    {"asset_id": "p-100", "description": "Plant", "plant": "1000", "asset_type": "PLANT"},
    {"asset_id": "p-100-pmp1", "description": "Pump", "parent_id": "p-100", "plant": "1000",
     "asset_type": "PUMP", "criticality": "high", "manufacturer": "Acme"},
]


def test_all_templates_registered():
    assert {t.system for t in list_templates()} >= {"SAP", "IBM Maximo", "Infor EAM", "eMaint", "Maintenance Connection"}
    assert len(list_templates("sap")) == 5


def test_mapping_and_transforms():
    rows, report = validate_records(get_template("sap_pm_equipment"), GOOD)
    assert report.ok
    assert rows[1]["EQUNR"] == "P-100-PMP1" and rows[1]["ABCKZ"] == "A" and rows[1]["TPLNR"] == "P-100"


def test_validation_errors():
    bad = [{"asset_id": "X" * 30, "description": "d" * 50, "asset_type": "T", "plant": "1000"},
           {"asset_id": "X" * 30, "description": "ok", "plant": "1000"}]
    _, report = validate_records(get_template("sap_pm_equipment"), bad)
    msgs = [i.message for i in report.errors]
    assert any("exceeds max 18" in m for m in msgs)
    assert any("exceeds max 40" in m for m in msgs)
    assert any("required" in m for m in msgs)
    assert any("duplicate" in m for m in msgs)


def test_bom_duplicate_key_includes_item():
    recs = [{"asset_id": "E1", "item_number": "10", "part_number": "m1", "quantity": 1, "plant": "1000"},
            {"asset_id": "E1", "item_number": "20", "part_number": "m2", "quantity": 2, "plant": "1000"}]
    assert validate_records(get_template("sap_pm_bom"), recs)[1].ok


def test_self_parent_error():
    _, report = validate_records(get_template("maximo_asset"), [{"asset_id": "A", "parent_id": "a", "description": "x", "site": "S"}])
    assert any("own parent" in i.message for i in report.errors)


def test_excel_export_and_audit(tmp_path):
    audit = AuditTrail(tmp_path / "audit.jsonl", user="tester")
    res = export_records("sap_pm_equipment", GOOD, tmp_path / "eq.xlsx", audit=audit)
    wb = load_workbook(res.path)
    assert wb.sheetnames == ["SAP PM - Equipment", "Field Guide"]
    ws = wb.active
    assert [c.value for c in ws[1]][:2] == ["EQUNR", "EQKTX"] and ws.max_row == 3
    (entry,) = audit.entries()
    assert entry["action"] == "export" and entry["rows"] == 2 and len(entry["sha256"]) == 64


def test_blocked_export_audited_and_no_file(tmp_path):
    audit = AuditTrail(tmp_path / "audit.jsonl")
    with pytest.raises(ExportBlocked):
        export_records("sap_mm_material", [{"part_number": "x"}], tmp_path / "m.xlsx", audit=audit)
    assert not (tmp_path / "m.xlsx").exists()
    assert audit.entries()[0]["action"] == "export_blocked"


def test_allow_errors_adds_validation_sheet(tmp_path):
    res = export_records("sap_mm_material", [{"part_number": "x"}], tmp_path / "m.xlsx", allow_errors=True)
    assert "Validation" in load_workbook(res.path).sheetnames


def test_csv_and_formula_guard(tmp_path):
    recs = [{"asset_id": "a1", "description": "=cmd()", "site": "S"}]
    res = export_records("maximo_asset", recs, tmp_path / "a.csv")
    assert "'=cmd()" in res.path.read_text()
