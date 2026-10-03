import io

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from asset_hierarchy.web.app import create_app


def client(tmp_path):
    return TestClient(create_app(str(tmp_path / "d.sqlite"), str(tmp_path / "a.jsonl"), str(tmp_path / "x")))


def test_full_flow(tmp_path):
    c = client(tmp_path)
    assert "Asset Hierarchy" in c.get("/").text
    pid = c.post("/api/projects", json={"site": "ACME", "prefix": "AP", "sap_maint_plant": "1000"}).json()["project"]["id"]
    prod = c.post(f"/api/projects/{pid}/nodes", json={"is_area": True, "description": "Production"}).json()["nodes"][0]
    d = c.post(f"/api/projects/{pid}/nodes", json={"parent_id": prod["id"], "description": "pump centrifugal",
                                                  "fields": {"tag_number": "P-1", "manufacturer": "Grundfos"}}).json()
    pump = [n for n in d["nodes"] if not n["is_area"]][0]
    assert pump["sap_name"] and pump["parent_id"] == prod["id"]
    assert c.post(f"/api/projects/{pid}/nodes/{prod['id']}/move", json={"parent_id": pump["id"]}).status_code == 400  # FLOC under equipment
    assert c.get(f"/api/projects/{pid}/preflight/sap_pm_equipment").status_code == 200
    r = c.get(f"/api/projects/{pid}/export/lsmw")
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content))
    assert "New Equipment" in wb.sheetnames
    # round-trip import into a fresh project
    pid2 = c.post("/api/projects", json={"site": "B", "prefix": "AP"}).json()["project"]["id"]
    imp = c.post(f"/api/projects/{pid2}/import", files={"file": ("l.xlsx", r.content)}).json()
    assert imp["stats"]["added"] == 2 and len(imp["nodes"]) == 2
    assert c.delete(f"/api/projects/{pid}/nodes/{pump['id']}").status_code == 200
    assert {e["action"] for e in c.get("/api/audit").json()} >= {"export", "import"}
    assert c.get("/api/projects/nope").status_code == 404
