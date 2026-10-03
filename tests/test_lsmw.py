import pytest
from openpyxl import Workbook, load_workbook

from asset_hierarchy import Hierarchy, Project
from asset_hierarchy.lsmw import EQ_IDX, FL_IDX, export_lsmw
from asset_hierarchy.sapimport import apply_sap_import, parse_sap_template


def build():
    p = Project(site="ACME", prefix="AP", sap_maint_plant="1000", sap_planning_plant="1000", sap_company_code="1000",
                sap_cost_center="1000100", sap_work_center="MMECH", sap_planner_group="PL2")
    h = Hierarchy(p)
    prod = h.add(is_area=True, description="Production")
    pent = h.add(parent_id=prod.id, is_area=True, description="Treater 2 - Penthouse")
    pump = h.add(parent_id=pent.id, description="pump-centrifugal-cw", tag_number="P-101", manufacturer="Grundfos",
                 model="CR 15-4", year="2018", catalog_code="ZCM0017")
    h.add(parent_id=pump.id, description="motor-electric", tag_number="M-101", manufacturer="ABB", catalog_code="ZCM0065")
    return h


def rows(ws):
    return [[c.value for c in r] for r in ws.iter_rows()]


def test_layout_matches_client_template(tmp_path):
    r = export_lsmw(build(), tmp_path / "l.xlsx", today="20260101")
    wb = load_workbook(r["path"])
    assert wb.sheetnames == ["New Functional Locations", "Change Existing Functional Loc", "New Equipment",
                             "Change Existing Equipment", "To be deleted", "Validation Summary"]
    fl, eq = rows(wb["New Functional Locations"]), rows(wb["New Equipment"])
    assert fl[5][0] == "ACME" and fl[6][0] == "Field" and fl[6][1] == "TPLNR" and fl[7][1] == 30      # rows 6,7,8
    assert fl[13][0] == "Example" and fl[16][1] == "AP-PROD" and fl[16][FL_IDX["TPLMA"]] is None        # data on row 17
    assert fl[17][1] == "AP-PROD-PENT" and fl[17][FL_IDX["TPLMA"]] == "AP-PROD"
    assert fl[17][39] == "AREA" and fl[17][36] == "20260101" and fl[17][FL_IDX["ARBPL"]] == "MMECH"
    assert eq[6][EQ_IDX["HEQUI"]] == "HEQUI" and eq[6][EQ_IDX["RBNR"]] == "RBNR"
    pump, motor = eq[16], eq[17]
    assert pump[EQ_IDX["EQKTX"]] == "PMP CENT CW GRUNDFOS CR 15-4" and pump[7] == "PUMP_CENTR" and pump[EQ_IDX["RBNR"]] == "ZCM0017"
    assert pump[EQ_IDX["TPLNR"]] == "AP-PROD-PENT" and pump[EQ_IDX["HEQUI"]] is None and pump[EQ_IDX["BAUJJ"]] == "2018"
    assert motor[EQ_IDX["HEQUI"]] == pump[EQ_IDX["LEGACYKEY"]]        # sub-equipment points at parent asset id
    assert r["violations"] == [] and wb["Validation Summary"]["A1"].value.startswith("All fields")


def test_truncation_and_overflow_reporting(tmp_path):
    h = build()
    next(n for n in h.nodes.values() if n.tag_number == "P-101").manufacturer = "M" * 45
    r = export_lsmw(h, tmp_path / "t.xlsx")
    assert any(t["field"] == "HERST" and t["limit"] == 30 for t in r["truncations"])
    assert "truncated" in load_workbook(r["path"])["Validation Summary"]["A1"].value
    h.project.prefix = "ZZZZZ"
    parent = next(n for n in h.nodes.values() if n.floc_code == "PENT")
    for _ in range(3):
        parent = h.add(parent_id=parent.id, is_area=True, description="x", floc_code="ABCDE")
    r = export_lsmw(h, tmp_path / "t2.xlsx")
    assert any("exceeds 30" in w for w in r["warnings"])


def test_modes(tmp_path):
    h = build()
    h.project.sap_structure = "one-to-one"
    r = export_lsmw(h, tmp_path / "o.xlsx")
    assert (r["flocs"], r["equipment"]) == (4, 2)
    h.project.sap_structure = "structure-leaf"
    r = export_lsmw(h, tmp_path / "s.xlsx")
    assert (r["flocs"], r["equipment"]) == (3, 1)        # leaf motor on FLOC sheet, parent pump on EQ sheet


def test_roundtrip_import(tmp_path):
    path = export_lsmw(build(), tmp_path / "rt.xlsx")["path"]
    parsed = parse_sap_template(path)
    assert not parsed.errors and not parsed.record_errors
    assert (len(parsed.flocs), len(parsed.equip)) == (2, 2)          # Example row skipped
    dst = Hierarchy(Project(site="NEW"))
    stats = apply_sap_import(dst, parsed, "new")
    assert stats["added"] == 4 and dst.project.prefix == "AP" and dst.project.sap_maint_plant == "1000"
    pump = next(n for n in dst.nodes.values() if n.tag_number == "P-101")
    motor = next(n for n in dst.nodes.values() if n.tag_number == "M-101")
    assert dst.floc_label(pump) == "AP-PROD-PENT" and motor.parent_id == pump.id and pump.year == "2018"
    again = apply_sap_import(dst, parsed, "merge")
    assert again["added"] == 0 and len(dst.nodes) == 4


def test_import_rejects_non_template(tmp_path):
    wb = Workbook(); wb.active.title = "Other"; wb.save(tmp_path / "x.xlsx")
    parsed = parse_sap_template(tmp_path / "x.xlsx")
    assert parsed.errors
    with pytest.raises(ValueError):
        apply_sap_import(Hierarchy(Project(site="s")), parsed)
